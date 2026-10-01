"""Tashqi xizmatlar: AI (OpenAI-mos API), rasm generatsiya, yt-dlp, AudD, docx/pptx."""
import asyncio
import base64
import json
import re
import tempfile
from io import BytesIO
from urllib.parse import quote

import aiohttp

import config as cfg

TMP = tempfile.gettempdir()
YT_SEM = asyncio.Semaphore(3)  # bir vaqtda ko'pi bilan 3 ta yuklash


class NoKey(Exception):
    pass


class AIError(Exception):
    pass


async def _cfg(db, scope):
    key = await db.get_setting(scope, "api_key") or cfg.OPENAI_API_KEY
    base = (await db.get_setting(scope, "base_url") or cfg.OPENAI_BASE_URL).rstrip("/")
    model = await db.get_setting(scope, "model") or cfg.OPENAI_MODEL
    return key, base, model


async def ask_ai(db, scope, messages, max_tokens=2500) -> str:
    key, base, model = await _cfg(db, scope)
    if not key:
        raise NoKey()
    async with aiohttp.ClientSession(timeout=aiohttp.ClientTimeout(total=150)) as s:
        async with s.post(f"{base}/chat/completions", headers={"Authorization": f"Bearer {key}"},
                          json={"model": model, "messages": messages, "max_tokens": max_tokens}) as r:
            d = await r.json(content_type=None)
            if r.status != 200:
                raise AIError((d.get("error") or {}).get("message", f"HTTP {r.status}") if isinstance(d, dict) else r.status)
            return d["choices"][0]["message"]["content"].strip()


async def generate_image(db, scope, prompt) -> bytes:
    """Kalit bo'lsa OpenAI-mos API, bo'lmasa (yoki xato bo'lsa) bepul Pollinations."""
    key, base, _ = await _cfg(db, scope)
    async with aiohttp.ClientSession(timeout=aiohttp.ClientTimeout(total=150)) as s:
        if key:
            try:
                async with s.post(f"{base}/images/generations", headers={"Authorization": f"Bearer {key}"},
                                  json={"model": cfg.IMAGE_MODEL, "prompt": prompt, "n": 1, "size": "1024x1024"}) as r:
                    d = await r.json(content_type=None)
                    if r.status == 200:
                        it = d["data"][0]
                        if it.get("b64_json"):
                            return base64.b64decode(it["b64_json"])
                        async with s.get(it["url"]) as g:
                            return await g.read()
            except Exception:
                pass
        async with s.get("https://image.pollinations.ai/prompt/" + quote(prompt),
                         params={"width": 1024, "height": 1024, "nologo": "true"}) as r:
            if r.status != 200:
                raise AIError("Rasm xizmati javob bermadi, keyinroq urinib ko'ring")
            return await r.read()


def extract_json(t: str):
    t = re.sub(r"```(?:json)?", "", t)
    starts = [x for x in (t.find("["), t.find("{")) if x >= 0]
    return json.loads(t[min(starts):max(t.rfind("]"), t.rfind("}")) + 1])


# ---------- yt-dlp (bloklovchi — to_thread orqali chaqiriladi) ----------
def _opts(**extra):
    o = {"quiet": True, "noplaylist": True, "outtmpl": f"{TMP}/%(id)s.%(ext)s", **extra}
    if cfg.YTDLP_COOKIES:
        o["cookiefile"] = cfg.YTDLP_COOKIES
    return o


def _yt_search(q):
    import yt_dlp
    with yt_dlp.YoutubeDL(_opts(extract_flat=True, skip_download=True)) as y:
        info = y.extract_info(f"ytsearch10:{q}", download=False)
    return [e for e in info.get("entries", []) if e and (e.get("duration") or 0) <= 1200][:8]


def _yt_download(url, audio=False):
    import yt_dlp
    fmt = ("bestaudio[ext=m4a]/bestaudio/best" if audio else
           "best[ext=mp4][filesize<48M]/best[ext=mp4]/best")
    with yt_dlp.YoutubeDL(_opts(format=fmt, max_filesize=49 * 1024 * 1024)) as y:
        info = y.extract_info(url, download=True)
        return y.prepare_filename(info), info


async def yt_search(q):
    async with YT_SEM:
        return await asyncio.to_thread(_yt_search, q)


async def yt_download(url, audio=False):
    async with YT_SEM:
        return await asyncio.to_thread(_yt_download, url, audio)


async def audd_recognize(token, data: bytes):
    """Audio parchasini AudD orqali aniqlaydi -> 'Ijrochi - Nom' yoki None."""
    form = aiohttp.FormData()
    form.add_field("api_token", token)
    form.add_field("file", data, filename="a.ogg")
    async with aiohttp.ClientSession(timeout=aiohttp.ClientTimeout(total=60)) as s:
        async with s.post("https://api.audd.io/", data=form) as r:
            d = await r.json(content_type=None)
    res = d.get("result")
    return f"{res['artist']} - {res['title']}" if res else None


# ---------- hujjat generatorlari ----------
def make_docx(title: str, text: str) -> bytes:
    from docx import Document
    d = Document()
    d.add_heading(title[:120], 0)
    for line in text.splitlines():
        s = line.strip().replace("**", "")
        if not s:
            continue
        if s.startswith("#"):
            d.add_heading(s.lstrip("# "), level=min(len(s) - len(s.lstrip("#")), 3))
        elif s.startswith(("- ", "* ")):
            d.add_paragraph(s[2:], style="List Bullet")
        else:
            d.add_paragraph(s)
    b = BytesIO()
    d.save(b)
    return b.getvalue()


def make_pptx(title: str, slides: list) -> bytes:
    from pptx import Presentation
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[0])
    s.shapes.title.text = title[:100]
    s.placeholders[1].text = "AI yordamida tayyorlandi"
    for sl in slides:
        s = p.slides.add_slide(p.slide_layouts[1])
        s.shapes.title.text = str(sl.get("title", ""))[:100]
        tf = s.placeholders[1].text_frame
        for i, b in enumerate(sl.get("bullets", [])[:6]):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.text = str(b)[:200]
    b = BytesIO()
    p.save(b)
    return b.getvalue()
